from docx import Document
import shutil
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def make_title(doc, title: str):
    return doc.add_heading(title, 0)

def make_score(doc, score: str):
    return doc.add_paragraph("score is " + score)

def make_info(doc, info: str):
    return doc.add_paragraph(info)

def make_div(doc):
    """
    设置分割线。效果等同于在 Word 中输入`---`后快捷插入的分割线。
    """
    p = doc.add_paragraph()
    # 获取段落属性
    p_pr = p._element.get_or_add_pPr()
    # 创建边框元素
    p_bdr = OxmlElement('w:pBdr')
    # 创建下边框元素
    bottom_border = OxmlElement('w:bottom')
    bottom_border.set(qn('w:val'), 'single')  # 边框样式为实线
    bottom_border.set(qn('w:sz'), '4')  # 边框宽度（4 表示 0.5 磅）
    bottom_border.set(qn('w:space'), '0')  # 边框与文本的间距
    bottom_border.set(qn('w:color'), '000000')  # 边框颜色（黑色）
    # 将下边框添加到边框元素
    p_bdr.append(bottom_border)

    # 将边框元素添加到段落属性
    p_pr.append(p_bdr)

    return p

# 复制文件
shutil.copyfile('./md4report/assets/md4report.docx', './md4report/assets/test.docx')

# 打开复制后的docx文件
doc = Document('./md4report/assets/test.docx')

starter = doc.paragraphs[0]  # 内容开头

meta = [
    make_title(doc, "Linux课程"),
    make_score(doc, "100"),
    make_info(doc, "This is CS50! The art of programming."),
    make_div(doc)
]

# 将元数据插入到文档开头，相关用法见“lxml”python库。
for m in meta:
    starter._p.addprevious(m._p)

# 保存修改后的文档
doc.save('./md4report/assets/test.docx')
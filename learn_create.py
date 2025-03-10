from docx import Document
# from docx.document import Document as TDocument
from docx.oxml.xmlchemy import OxmlElement
from docx.table import _Cell, Table
# from docx.text.paragraph import Paragraph, Run
# from docx.image.image import Image
# from docx.enum.text import WD_BREAK
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_school_name(doc):
    doc.add_heading('广州航海学院', 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading('____实验报告', 1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    
def add_student_info(doc):
    meta_data = ["专业班级","计院","实验日期","2025.01.01","姓名","派蒙","学号","20241111","实验名称","清华拳实战","指导老师","辅导员"]
    table: Table = doc.add_table(rows=3, cols=4)

    for i in range(3):
        for j in range(4):
            # 计算当前项在列表中的索引
            index = i * 4 + j
            # 将列表中的项填入表格单元格
            table.cell(i, j).text = meta_data[index]

    # 遍历表格的每一行
    for row in table.rows:
        # 遍历每一行的单元格
        for idx, cell in enumerate(row.cells):
            # 获取单元格的边框属性
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()

            # 创建边框元素
            tcBorders = OxmlElement('w:tcBorders')

            # 如果是第2列或第4列，设置下边框
            if idx == 1 or idx == 3:
                cell.width = Inches(3)
                bottom = OxmlElement('w:bottom')
                bottom.set(qn('w:val'), 'single')
                bottom.set(qn('w:sz'), '4')
                bottom.set(qn('w:space'), '0')
                bottom.set(qn('w:color'), '000000')
                tcBorders.append(bottom)
            else:
                # 如果是第1列或第3列，移除所有边框
                cell.width = Inches(2)

                for border_name in ['top', 'left', 'bottom', 'right']:
                    border = OxmlElement(f'w:{border_name}')
                    border.set(qn('w:val'), 'nil')
                    tcBorders.append(border)

            # 将边框元素添加到单元格属性中
            tcPr.append(tcBorders)
            # 如果是第1列或第3列，设置分散对齐
            if idx == 0 or idx == 2:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.DISTRIBUTE
            if idx == 1 or idx == 3:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc = Document()
add_school_name(doc)
add_student_info(doc)
doc.save("info.docx")
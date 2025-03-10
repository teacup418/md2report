from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

# 创建一个新的文档
# doc = Document()
doc = Document('./md4report/assets/test.docx')
# 设置页面宽度和页边距
section = doc.sections[0]
section.page_width = Inches(8.5)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

# 添加一个3行4列的表格
table = doc.add_table(rows=3, cols=4)

# 遍历表格的每一行
for row in table.rows:
    # 遍历每一行的单元格
    for idx, cell in enumerate(row.cells):
        # 获取单元格的边框属性
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()

        # 创建边框元素
        tcBorders = OxmlElement('w:tcBorders')

        # 如果是第2列或第4列，设置下边框
        if idx == 1 or idx == 3:
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '4')
            bottom.set(qn('w:space'), '0')
            bottom.set(qn('w:color'), '000000')
            tcBorders.append(bottom)
        else:
            # 如果是第1列或第3列，移除所有边框
            for border_name in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement(f'w:{border_name}')
                border.set(qn('w:val'), 'nil')
                tcBorders.append(border)

        # 将边框元素添加到单元格属性中
        tcPr.append(tcBorders)

# 保存文档
doc.save('./md4report/assets/test.docx')